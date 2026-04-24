import os
import queue
import threading
import tkinter as tk
from tkinter import filedialog, messagebox

import speech_recognition as sr
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


class VoiceTranscriberApp:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title("Транскрибатор голоса")
        self.root.geometry("900x600")
        self.root.configure(bg="#111827")

        self.recognizer = sr.Recognizer()
        self.is_recording = False
        self.audio_queue: queue.Queue[sr.AudioData] = queue.Queue()
        self.transcript_parts: list[str] = []
        self.stop_listening = None

        self._build_ui()

    def _build_ui(self) -> None:
        top_frame = tk.Frame(self.root, bg="#111827")
        top_frame.pack(fill="x", padx=24, pady=(24, 12))

        title = tk.Label(
            top_frame,
            text="Голосовой транскрибатор",
            font=("Segoe UI", 22, "bold"),
            bg="#111827",
            fg="#F3F4F6",
        )
        title.pack(anchor="w")

        subtitle = tk.Label(
            top_frame,
            text="Нажмите большую кнопку Старт/Стоп для записи и распознавания речи",
            font=("Segoe UI", 11),
            bg="#111827",
            fg="#9CA3AF",
        )
        subtitle.pack(anchor="w", pady=(6, 0))

        control_frame = tk.Frame(self.root, bg="#111827")
        control_frame.pack(fill="x", padx=24, pady=10)

        self.toggle_button = tk.Button(
            control_frame,
            text="СТАРТ",
            command=self.toggle_recording,
            font=("Segoe UI", 26, "bold"),
            width=16,
            height=2,
            bg="#10B981",
            fg="#F9FAFB",
            activebackground="#059669",
            activeforeground="#FFFFFF",
            bd=0,
            relief="flat",
            cursor="hand2",
        )
        self.toggle_button.pack(pady=12)

        save_frame = tk.Frame(self.root, bg="#111827")
        save_frame.pack(fill="x", padx=24, pady=(0, 16))

        btn_style = {
            "font": ("Segoe UI", 11, "bold"),
            "bg": "#374151",
            "fg": "#F9FAFB",
            "activebackground": "#4B5563",
            "activeforeground": "#FFFFFF",
            "bd": 0,
            "relief": "flat",
            "cursor": "hand2",
            "padx": 16,
            "pady": 8,
        }

        tk.Button(save_frame, text="Сохранить TXT", command=self.save_txt, **btn_style).pack(
            side="left", padx=(0, 10)
        )
        tk.Button(save_frame, text="Сохранить DOCX", command=self.save_docx, **btn_style).pack(
            side="left", padx=10
        )
        tk.Button(save_frame, text="Сохранить PDF", command=self.save_pdf, **btn_style).pack(
            side="left", padx=10
        )

        self.status_label = tk.Label(
            self.root,
            text="Статус: ожидание",
            font=("Segoe UI", 11),
            bg="#111827",
            fg="#D1D5DB",
        )
        self.status_label.pack(anchor="w", padx=24)

        text_frame = tk.Frame(self.root, bg="#111827")
        text_frame.pack(fill="both", expand=True, padx=24, pady=(12, 24))

        self.text_widget = tk.Text(
            text_frame,
            wrap="word",
            font=("Consolas", 12),
            bg="#1F2937",
            fg="#F3F4F6",
            insertbackground="#F3F4F6",
            selectbackground="#2563EB",
            bd=0,
            padx=12,
            pady=12,
        )
        self.text_widget.pack(fill="both", expand=True)

    def toggle_recording(self) -> None:
        if not self.is_recording:
            self.start_recording()
        else:
            self.stop_recording()

    def start_recording(self) -> None:
        try:
            with sr.Microphone() as source:
                self.recognizer.adjust_for_ambient_noise(source, duration=0.8)

            self.stop_listening = self.recognizer.listen_in_background(
                sr.Microphone(),
                self._audio_callback,
                phrase_time_limit=8,
            )
            self.is_recording = True
            self.toggle_button.configure(text="СТОП", bg="#DC2626", activebackground="#B91C1C")
            self.status_label.configure(text="Статус: запись и распознавание...")

            threading.Thread(target=self._process_audio_queue, daemon=True).start()
        except Exception as exc:
            messagebox.showerror("Ошибка", f"Не удалось открыть микрофон: {exc}")

    def stop_recording(self) -> None:
        if self.stop_listening:
            self.stop_listening(wait_for_stop=False)
            self.stop_listening = None

        self.is_recording = False
        self.toggle_button.configure(text="СТАРТ", bg="#10B981", activebackground="#059669")
        self.status_label.configure(text="Статус: запись остановлена")

    def _audio_callback(self, recognizer: sr.Recognizer, audio: sr.AudioData) -> None:
        self.audio_queue.put(audio)

    def _process_audio_queue(self) -> None:
        while self.is_recording or not self.audio_queue.empty():
            try:
                audio = self.audio_queue.get(timeout=0.5)
            except queue.Empty:
                continue

            try:
                text = self.recognizer.recognize_google(audio, language="ru-RU")
                if text.strip():
                    self.transcript_parts.append(text)
                    self.root.after(0, self._append_text, text)
            except sr.UnknownValueError:
                continue
            except sr.RequestError as exc:
                self.root.after(
                    0,
                    lambda: self.status_label.configure(
                        text=f"Статус: ошибка сервиса распознавания ({exc})"
                    ),
                )
                break

    def _append_text(self, text: str) -> None:
        current = self.text_widget.get("1.0", "end-1c").strip()
        delimiter = "\n" if current else ""
        self.text_widget.insert("end", f"{delimiter}{text}")
        self.text_widget.see("end")

    def _get_transcript(self) -> str:
        return self.text_widget.get("1.0", "end-1c").strip()

    def save_txt(self) -> None:
        content = self._get_transcript()
        if not content:
            messagebox.showwarning("Пусто", "Нет текста для сохранения.")
            return

        filepath = filedialog.asksaveasfilename(
            defaultextension=".txt", filetypes=[("Text files", "*.txt")]
        )
        if not filepath:
            return

        with open(filepath, "w", encoding="utf-8") as f:
            f.write(content)
        messagebox.showinfo("Готово", "TXT файл сохранен.")

    def save_docx(self) -> None:
        content = self._get_transcript()
        if not content:
            messagebox.showwarning("Пусто", "Нет текста для сохранения.")
            return

        filepath = filedialog.asksaveasfilename(
            defaultextension=".docx", filetypes=[("Word files", "*.docx")]
        )
        if not filepath:
            return

        doc = Document()
        doc.add_heading("Транскрипция", level=1)
        doc.add_paragraph(content)
        doc.save(filepath)
        messagebox.showinfo("Готово", "DOCX файл сохранен.")

    def save_pdf(self) -> None:
        content = self._get_transcript()
        if not content:
            messagebox.showwarning("Пусто", "Нет текста для сохранения.")
            return

        filepath = filedialog.asksaveasfilename(
            defaultextension=".pdf", filetypes=[("PDF files", "*.pdf")]
        )
        if not filepath:
            return

        font_name = "Helvetica"
        font_candidates = [
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "C:/Windows/Fonts/arial.ttf",
            "/Library/Fonts/Arial Unicode.ttf",
        ]
        for font_path in font_candidates:
            if os.path.exists(font_path):
                pdfmetrics.registerFont(TTFont("AppFont", font_path))
                font_name = "AppFont"
                break

        c = canvas.Canvas(filepath, pagesize=A4)
        width, height = A4
        c.setFont(font_name, 12)

        x = 40
        y = height - 40
        for paragraph in content.split("\n"):
            words = paragraph.split()
            line = ""
            for word in words:
                test_line = f"{line} {word}".strip()
                if c.stringWidth(test_line, font_name, 12) < width - 80:
                    line = test_line
                else:
                    c.drawString(x, y, line)
                    y -= 18
                    line = word
                    if y < 40:
                        c.showPage()
                        c.setFont(font_name, 12)
                        y = height - 40
            if line:
                c.drawString(x, y, line)
                y -= 20
                if y < 40:
                    c.showPage()
                    c.setFont(font_name, 12)
                    y = height - 40

        c.save()
        messagebox.showinfo("Готово", "PDF файл сохранен.")


if __name__ == "__main__":
    root = tk.Tk()
    app = VoiceTranscriberApp(root)
    root.mainloop()
